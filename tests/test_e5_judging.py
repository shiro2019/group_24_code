"""Regression checks for E5 package safety, legacy input and scoring."""
import csv
import io
import json
import random
import shutil
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from PIL import Image

from src import e5_judging as e5


class JudgingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.prompts = [{"prompt_id": "p1", "prompt": "two red balls", "challenge": "quantity"}]
        self.prompt_path = self.root / "prompts.jsonl"
        self.prompt_path.write_text(json.dumps(self.prompts[0]) + "\n")
        self.rows = e5.plan(self.prompts, random.Random(2026))
        for row in self.rows:
            row["config"] = "frozen"
        self.key = self.root / "key.csv"
        self.write_csv(self.key, self.rows)
        self.sheets = self.root / "sheets"
        self.sheets.mkdir()
        self.args = SimpleNamespace(key=self.key, sheets=self.sheets, out=self.root / "results",
                                    resolved=None, images=None, packages=None, seed=2026,
                                    prompts=self.prompt_path)

    def write_csv(self, path, rows):
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)

    def workbook(self, language="en", primary=None, recheck=None):
        primary = ["pass"] * 4 if primary is None else primary
        recheck = primary if recheck is None else recheck
        tabs = ("Primary", "Recheck") if language == "en" else ("初判", "复核")
        labels = {"pass": "通过", "fail": "失败", "unsure": "无法判断", "": ""}
        wb = Workbook()
        wb.remove(wb.active)
        for tab, values in zip(tabs, (primary, recheck)):
            ws = wb.create_sheet(tab)
            ws.append(e5.HEADER)
            for i, (row, verdict) in enumerate(zip(self.rows, values), 1):
                ws.append([i, row["image_id"], None, "Quantity", row["prompt"], "rule",
                           labels[verdict] if language == "zh" else verdict, 123])
        path = self.sheets / "returned.xlsx"
        wb.save(path)
        return path

    def score(self):
        with redirect_stdout(io.StringIO()):
            e5.score(self.args)
        with (self.args.out / "e5_failure.csv").open() as f:
            return list(csv.DictReader(f))

    def test_legacy_and_english_give_identical_results(self):
        self.workbook(primary=["fail", "pass", "unsure", "pass"])
        english = self.score()
        self.workbook("zh", primary=["fail", "pass", "unsure", "pass"])
        self.assertEqual(self.score(), english)
        unsure = next(r for r in english if r["method"] == "SDXL-Turbo" and r["challenge"] == "all")
        self.assertEqual(unsure["unsure"], "1")
        self.assertEqual(unsure["fail_rate"], "")
        self.assertEqual(unsure["status"], "final")

    def test_duplicate_sheet_and_invalid_verdict_rejected(self):
        path = self.workbook()
        duplicate = self.sheets / "copy.xlsx"
        shutil.copy(path, duplicate)
        with self.assertRaisesRegex(ValueError, "duplicate judgment"):
            e5.read_sheets(self.sheets)
        duplicate.unlink()
        self.workbook(primary=["maybe"] * 4)
        with self.assertRaisesRegex(ValueError, "invalid verdict"):
            e5.read_sheets(self.sheets)

    def test_lock_file_is_ignored_and_notes_may_be_numbers(self):
        self.workbook()
        (self.sheets / "~$returned.xlsx").write_text("not a workbook")
        self.assertEqual(len(e5.read_sheets(self.sheets)), 8)
        self.assertEqual(next(iter(e5.read_sheets(self.sheets).values()))[1], "123")

    def test_incomplete_rechecks_are_missing_not_disagreements(self):
        self.workbook(recheck=[""] * 4)
        rows = self.score()
        self.assertTrue(all(r["status"] == "PROVISIONAL" for r in rows))
        summary = (self.args.out / "e5_summary.txt").read_text()
        self.assertIn("missing or invalid: 4", summary)
        self.assertIn("disagreements: 0", summary)

    def test_chinese_final_decision_and_overwrite_guard(self):
        self.workbook(primary=["pass"] * 4, recheck=["fail", "pass", "pass", "pass"])
        self.score()
        path = self.args.out / "disagreements.csv"
        with path.open(encoding="utf-8-sig") as f:
            decisions = list(csv.DictReader(f))
        decisions[0].update(primary="通过", recheck="失败", final="失败")
        self.write_csv(path, decisions)
        before = path.read_bytes()
        with self.assertRaisesRegex(ValueError, "--resolved"):
            self.score()
        self.assertEqual(path.read_bytes(), before)
        self.args.resolved = path
        rows = self.score()
        self.assertEqual(rows[0]["fail_rate"], "1.000")
        self.assertEqual(rows[0]["status"], "final")
        self.workbook(primary=["unsure", "pass", "pass", "pass"], recheck=["fail", "pass", "pass", "pass"])
        with self.assertRaisesRegex(ValueError, "outdated"):
            self.score()

    def test_nonexistent_resolution_file_is_not_ignored(self):
        self.workbook()
        self.args.resolved = self.root / "absent.csv"
        with self.assertRaises(FileNotFoundError):
            self.score()
        self.assertFalse(self.args.out.exists())

    def test_unknown_image_and_incomplete_key_rejected(self):
        path = self.workbook()
        wb = load_workbook(path)
        wb["Primary"]["B2"] = "unknown"
        wb.save(path)
        wb.close()
        with self.assertRaisesRegex(ValueError, "unexpected image"):
            self.score()
        self.write_csv(self.key, self.rows[:-1])
        with self.assertRaisesRegex(ValueError, "exactly one image"):
            self.score()

    def test_pack_never_removes_existing_directory(self):
        sentinel = self.root / "important.txt"
        sentinel.write_text("keep")
        with self.assertRaises(FileExistsError):
            e5.pack(SimpleNamespace(out=self.root))
        self.assertEqual(sentinel.read_text(), "keep")

    def test_pack_english_workbooks_and_fixed_image_size(self):
        images = self.root / "images"
        for method in e5.METHODS:
            folder = images / method / "frozen"
            folder.mkdir(parents=True)
            Image.new("RGB", (32, 32), "red").save(folder / "p1_0.png")
        out = self.root / "packages"
        with redirect_stdout(io.StringIO()):
            e5.pack(SimpleNamespace(out=out, images=images, prompts=self.prompt_path, seed=2026))
        count = 0
        for folder in out.glob("E5_judge_*"):
            count += len(e5.read_sheets(folder))
            wb = load_workbook(next(folder.glob("*.xlsx")))
            self.assertEqual(wb.sheetnames, ["Guide", "Primary", "Recheck"])
            wb.close()
            for path in folder.glob("images/*.png"):
                with Image.open(path) as image:
                    self.assertEqual(image.size, (512, 512))
        self.assertEqual(count, 8)

    def test_single_failure_example_and_no_examples_after_rerun(self):
        packages = self.root / "packages" / "E5判定_A" / "images"
        packages.mkdir(parents=True)
        for row in self.rows:
            Image.new("RGB", (32, 32), "blue").save(packages / (row["image_id"] + ".png"))
        self.args.packages = packages.parent.parent
        self.workbook(primary=["fail", "pass", "pass", "pass"])
        self.score()
        self.assertTrue((self.args.out / "e5_failure_cases.png").is_file())
        self.workbook()
        self.score()
        self.assertFalse((self.args.out / "e5_failure_cases.png").exists())

    def test_metrics_use_frozen_config_and_require_complete_scores(self):
        for method in e5.METHODS:
            for config in ("frozen", "unrelated"):
                folder = self.root / method / config
                folder.mkdir(parents=True)
                (folder / "per_image_scores.jsonl").write_text(json.dumps({
                    "prompt_id": "p1", "seed": 0, "clip_score": 20 if config == "frozen" else 90,
                    "image_reward": 1}) + "\n")
        metrics = e5.per_challenge_metrics(self.root, self.rows)
        self.assertEqual(metrics[("sdxl_base", "all")]["clip_score"], 20)
        path = self.root / "sdxl_base" / "frozen" / "per_image_scores.jsonl"
        path.write_text("")
        with self.assertRaisesRegex(ValueError, "missing scores"):
            e5.per_challenge_metrics(self.root, self.rows)

    def test_degenerate_kappa_and_known_wilson_interval(self):
        self.assertIsNone(e5.cohen_kappa([("pass", "pass")] * 4))
        self.assertAlmostEqual(e5.cohen_kappa([("pass", "pass"), ("fail", "fail")]), 1)
        low, high = e5.wilson(5, 10)
        self.assertAlmostEqual(low, 0.23659, places=4)
        self.assertAlmostEqual(high, 0.76341, places=4)
        self.assertEqual(e5.wilson(0, 0), (None, None))


if __name__ == "__main__":
    unittest.main()

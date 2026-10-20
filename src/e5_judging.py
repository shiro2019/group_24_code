"""E5 human judging: blind packages for the judges, then failure rates from their sheets.

    python -m src.e5_judging pack  --images outputs/e5/parti --out outputs/e5_judging
    python -m src.e5_judging score --key outputs/e5_judging/_key/key.csv --sheets returned/ --out outputs/e5_results

pack: every E5 image (4 methods x 87 frozen prompts) gets a random image id. Prompts are
dealt to the four judges per challenge, so each judge sees all four images of their
prompts and a similar mix of challenges. 20% of the images, sampled per method and
challenge, also go to a second judge (the next one in JUDGES) for an independent check.
Each judge gets a folder with the images (all at 512 px, re-saved without metadata) and a sheet with
two tabs, primary and re-check. Method names appear only in _key/key.csv.

score: reads the returned sheets and writes the re-check disagreements for discussion,
the failure rate per method and challenge with a 95% Wilson interval, re-check
agreement and Cohen's kappa, and the figures. A filled-in "final" column in the
disagreements file overrides the primary judgment after the discussion; until every
judgment is in and every disagreement is resolved the results are marked PROVISIONAL.
"""

import argparse
import csv
import json
import math
import random
import shutil
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from src.common import REPO_ROOT

PROMPTS = REPO_ROOT / "configs" / "e5" / "parti_prompts.jsonl"
JUDGES = ["A", "B", "C", "D"]
RECHECK_SHARE = 0.2
SEED = 2026
METHODS = {"sdxl_base": "SDXL + DPM-Solver++", "lcm_lora": "LCM-LoRA",
           "sdxl_turbo": "SDXL-Turbo", "sdxl_lightning": "SDXL-Lightning"}
CHALLENGES = {  # challenge -> (label, what passes)
    "quantity": ("Quantity", "Every explicitly counted object has the exact requested count"),
    "writing": ("Writing", "All requested text is present and spelled correctly (ignore case and punctuation)"),
    "position": ("Position / properties", "All requested relative positions and properties are correct"),
}
VERDICTS = ["pass", "fail", "unsure"]
JUDGE_PX = 512  # every image is shown at this size: Turbo is native 512, the E5 configs save the rest at 1024
THUMB_PX = 256
FONT = "Arial"
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
HEADER = ["No.", "Image ID", "Image", "Challenge", "Prompt", "Pass criterion", "Verdict", "Notes"]


def load_prompts(path=PROMPTS):
    with open(path, encoding="utf-8") as f:
        rows = [json.loads(line) for line in f if line.strip()]
    ids = [r["prompt_id"] for r in rows]
    if not rows or len(ids) != len(set(ids)):
        raise ValueError("prompts must be non-empty with unique prompt IDs")
    if any(r["challenge"] not in CHALLENGES for r in rows):
        raise ValueError("unknown prompt challenge")
    return rows


def find_images(root, prompts):
    """{(method, prompt_id): path}; each method folder holds exactly one config."""
    found = {}
    for method in METHODS:
        configs = [d for d in (Path(root) / method).iterdir() if d.is_dir()]
        if len(configs) != 1:
            raise ValueError(f"{method}: expected one config folder, found {configs}")
        for p in prompts:
            path = configs[0] / f"{p['prompt_id']}_0.png"
            if not path.is_file():
                raise FileNotFoundError(f"missing {path}")
            found[(method, p["prompt_id"])] = (configs[0].name, path)
    return found


def plan(prompts, rng):
    """Image rows with id, judge and re-check judge."""
    judge_of, turn = {}, 0
    for challenge in CHALLENGES:
        ids = [p["prompt_id"] for p in prompts if p["challenge"] == challenge]
        rng.shuffle(ids)
        for pid in ids:
            judge_of[pid] = JUDGES[turn % len(JUDGES)]
            turn += 1
    rows = [{"prompt_id": p["prompt_id"], "prompt": p["prompt"], "challenge": p["challenge"],
             "method": m, "judge": judge_of[p["prompt_id"]], "recheck_judge": ""}
            for p in prompts for m in METHODS]
    for row, num in zip(rows, rng.sample(range(1000, 10000), len(rows))):
        row["image_id"] = f"img{num}"
    cells = defaultdict(list)
    for row in rows:
        cells[(row["method"], row["challenge"])].append(row)
    for cell in cells.values():
        for row in rng.sample(cell, math.ceil(RECHECK_SHARE * len(cell))):
            row["recheck_judge"] = JUDGES[(JUDGES.index(row["judge"]) + 1) % len(JUDGES)]
    return rows


def write_tab(ws, rows, folder):
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True)
    widths = [6, 11, 37, 22, 40, 30, 11, 30]
    for col, w in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, row in enumerate(rows, start=1):
        label, rule = CHALLENGES[row["challenge"]]
        ws.append([i, row["image_id"], None, label, row["prompt"], rule, None, None])
        r = ws.max_row
        thumb = folder / "thumbs" / f"{row['image_id']}.png"
        img = XLImage(str(thumb))
        img.anchor = f"C{r}"
        ws.add_image(img)
        ws.row_dimensions[r].height = THUMB_PX * 0.75 + 6
        for cell in ws[r]:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws[f"B{r}"].hyperlink = f"images/{row['image_id']}.png"
        for col in "GH":
            ws[f"{col}{r}"].fill = INPUT_FILL
        dv.add(f"G{r}")
    ws.freeze_panes = "A2"


def write_guide(ws, judge, n_primary, n_recheck):
    lines = [
        f"E5 human judging — Judge {judge}",
        "",
        "Fill only the yellow Verdict (dropdown) and Notes cells in Primary and Recheck.",
        "Click an image ID to open the 512 x 512 judging image in images/. Use it for text judgments.",
        "",
        "Judge only the specified capability, not image quality or aesthetics:",
        "Quantity: every explicitly counted object must have exactly the requested count. Missing or wrong objects fail; count occluded objects if distinguishable.",
        "Writing: all requested text must be present and spelled correctly (ignore case and punctuation). Extra unrelated text is allowed; missing or incorrect letters fail.",
        "Position / properties: all requested relative positions and properties (such as colors) must hold. Any missing requested object fails.",
        "Choose unsure when you cannot decide, and explain why in Notes.",
        "",
        "Image IDs are random and do not identify methods. Judge each of the four images for a prompt independently.",
        "Recheck contains another judge's images. Judge independently without discussion or consulting their sheet.",
        "Return this Excel file with its filename unchanged.",
        "",
        "Example (format only, not real data):",
    ]
    for line in lines:
        ws.append([line])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=5)
        ws.row_dimensions[ws.max_row].height = 45 if len(line) > 100 else 30
    ws.append(["Image ID", "Challenge", "Prompt", "Verdict", "Notes"])
    ws.append(["img1234", "Quantity", "three red balls on a table", "fail", "Only two balls"])
    ws.append([])
    ws.append(["Progress", "Filled", "Total"])
    for tab, count in (("Primary", n_primary), ("Recheck", n_recheck)):
        ws.append([tab, f"=COUNTA('{tab}'!G2:G{count + 1})" if count else 0, count])
    for col, width in zip("ABCDE", [15, 15, 40, 15, 30]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=cell.row == 1)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def pack(args):
    out = Path(args.out)
    if out.exists() or out.is_symlink():
        raise FileExistsError(f"{out} already exists; choose a new output directory")
    prompts = load_prompts(args.prompts)
    images = find_images(args.images, prompts)
    rng = random.Random(args.seed)
    rows = plan(prompts, rng)
    out.mkdir(parents=True, exist_ok=False)
    for row in rows:
        row["config"], src = images[(row["method"], row["prompt_id"])]
        for judge in {row["judge"], row["recheck_judge"]} - {""}:
            folder = out / f"E5_judge_{judge}"
            (folder / "images").mkdir(parents=True, exist_ok=True)
            (folder / "thumbs").mkdir(exist_ok=True)
            with Image.open(src) as im:  # re-save: drops PNG text chunks and other metadata
                clean = im.convert("RGB")
            if clean.size != (JUDGE_PX, JUDGE_PX):  # a size difference would give the method away
                clean = clean.resize((JUDGE_PX, JUDGE_PX), Image.BICUBIC)
            clean.save(folder / "images" / f"{row['image_id']}.png", format="PNG")
            clean.resize((THUMB_PX, THUMB_PX), Image.LANCZOS).save(
                folder / "thumbs" / f"{row['image_id']}.png", format="PNG")
    for judge in JUDGES:
        folder = out / f"E5_judge_{judge}"
        folder.mkdir(parents=True, exist_ok=True)
        mine = [r for r in rows if r["judge"] == judge]
        order = list(dict.fromkeys(r["prompt_id"] for r in mine))
        rng.shuffle(order)
        primary = []
        for pid in order:  # a prompt's four images together, in random method order
            group = [r for r in mine if r["prompt_id"] == pid]
            rng.shuffle(group)
            primary += group
        recheck = [r for r in rows if r["recheck_judge"] == judge]
        rng.shuffle(recheck)
        wb = Workbook()
        write_guide(wb.active, judge, len(primary), len(recheck))
        wb.active.title = "Guide"
        write_tab(wb.create_sheet("Primary"), primary, folder)
        write_tab(wb.create_sheet("Recheck"), recheck, folder)
        wb.calculation.fullCalcOnLoad = True  # the progress counts have no cached values
        wb.save(folder / f"E5_judging_{judge}.xlsx")
        if (folder / "thumbs").exists():
            shutil.rmtree(folder / "thumbs")  # embedded in the sheet
        print(f"{judge}: {len(primary)} primary ({len(order)} prompts), {len(recheck)} re-check")
    (out / "_key").mkdir()
    fields = ["image_id", "method", "config", "prompt_id", "challenge", "judge", "recheck_judge"]
    with open(out / "_key" / "key.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: r["image_id"]))
    print(f"{len(rows)} images, {sum(bool(r['recheck_judge']) for r in rows)} re-checked; key in {out / '_key'}")


# Input aliases keep already-distributed Chinese workbooks usable.
TAB_ALIASES = {"Primary": ("Primary", "初判"), "Recheck": ("Recheck", "复核")}
VERDICT_ALIASES = {"pass": "pass", "fail": "fail", "unsure": "unsure",
                   "通过": "pass", "失败": "fail", "无法判断": "unsure"}


def normalize_verdict(value):
    text = str(value).strip().lower() if value is not None else ""
    if not text:
        return ""
    if text not in VERDICT_ALIASES:
        raise ValueError(f"invalid verdict {value!r}; expected pass, fail or unsure")
    return VERDICT_ALIASES[text]


def read_sheets(folder):
    """Read English or legacy Chinese sheets; never silently overwrite a judgment."""
    folder = Path(folder)
    if not folder.is_dir():
        raise NotADirectoryError(folder)
    out = {}
    for path in sorted(folder.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue  # Excel lock file
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            for tab, aliases in TAB_ALIASES.items():
                matches = [name for name in aliases if name in wb.sheetnames]
                if len(matches) != 1:
                    raise ValueError(f"{path}: expected exactly one {tab} tab")
                for number, row in enumerate(wb[matches[0]].iter_rows(min_row=2, max_col=8, values_only=True), 2):
                    if not any(value is not None for value in row):
                        continue
                    if not row[1]:
                        raise ValueError(f"{path}:{tab}:{number}: missing image ID")
                    image_id = str(row[1]).strip()
                    key = (image_id, tab)
                    if key in out:
                        raise ValueError(f"{path}: duplicate judgment for {image_id} / {tab}")
                    try:
                        verdict = normalize_verdict(row[6])
                    except ValueError as exc:
                        raise ValueError(f"{path}:{tab}:{number}: {exc}") from exc
                    out[key] = (verdict, str(row[7]).strip() if row[7] is not None else "")
        finally:
            wb.close()
    return out


def wilson(k, n, z=1.96):
    """95% Wilson score interval for k failures out of n judged images."""
    if not n:
        return None, None
    p = k / n
    centre = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return max(0.0, centre - half), min(1.0, centre + half)


def cohen_kappa(pairs):
    """Agreement between primary and re-check verdicts beyond chance."""
    n = len(pairs)
    if not n:
        return None
    observed = sum(a == b for a, b in pairs) / n
    first, second = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    expected = sum(first[v] * second[v] for v in set(first) | set(second)) / (n * n)
    return (observed - expected) / (1 - expected) if expected < 1 else None


def per_challenge_metrics(root, key):
    """Mean CLIP Score and ImageReward per method and challenge, from the scorer's
    per_image_scores.jsonl in each E5 output folder."""
    sums = defaultdict(lambda: defaultdict(list))
    for method in METHODS:
        selected = [r for r in key if r["method"] == method]
        challenge = {r["prompt_id"]: r["challenge"] for r in selected}
        path = Path(root, method, selected[0]["config"], "per_image_scores.jsonl")
        seen = set()
        with open(path, encoding="utf-8") as f:
            for line in f:
                if not line.strip():
                    continue
                r = json.loads(line)
                pid = r["prompt_id"]
                if pid not in challenge or pid in seen or r.get("seed") != 0:
                    raise ValueError(f"{path}: unexpected or duplicate image: {pid}")
                seen.add(pid)
                for ch in (challenge[r["prompt_id"]], "all"):
                    for metric in ("clip_score", "image_reward"):
                        value = r.get(metric)
                        if not isinstance(value, (int, float)) or not math.isfinite(value):
                            raise ValueError(f"{path}: missing or invalid {metric} for {pid}")
                        sums[(method, ch)][metric].append(value)
        if seen != set(challenge):
            raise ValueError(f"{path}: missing scores for {sorted(set(challenge) - seen)}")
    return {k: {m: sum(v) / len(v) for m, v in d.items()} for k, d in sums.items()}


def score(args):
    with open(args.key, encoding="utf-8-sig") as f:
        key = list(csv.DictReader(f))
    prompts = load_prompts(args.prompts)
    prompt_map = {p["prompt_id"]: p for p in prompts}
    expected = {(m, p["prompt_id"]) for m in METHODS for p in prompts}
    actual = {(r["method"], r["prompt_id"]) for r in key}
    if actual != expected or len(key) != len(expected):
        raise ValueError("key must contain exactly one image per method and prompt")
    if len({r["image_id"] for r in key}) != len(key):
        raise ValueError("key contains duplicate image IDs")
    for r in key:
        if (r["challenge"] != prompt_map[r["prompt_id"]]["challenge"]
                or r["judge"] not in JUDGES
                or r["recheck_judge"] not in ["", *JUDGES]
                or r["recheck_judge"] == r["judge"]):
            raise ValueError(f"invalid key assignment for {r['image_id']}")
    if any(len({r["config"] for r in key if r["method"] == m}) != 1 for m in METHODS):
        raise ValueError("key must contain exactly one configuration per method")
    got = read_sheets(args.sheets)
    expected_slots = {(r["image_id"], "Primary") for r in key}
    expected_slots |= {(r["image_id"], "Recheck") for r in key if r["recheck_judge"]}
    if set(got) - expected_slots:
        raise ValueError(f"unexpected image IDs or tabs: {sorted(set(got) - expected_slots)[:5]}")
    out = Path(args.out)
    missing = [r["image_id"] for r in key if got.get((r["image_id"], "Primary"), ("",))[0] not in VERDICTS]
    missing += [r["image_id"] for r in key if r["recheck_judge"]
                and got.get((r["image_id"], "Recheck"), ("",))[0] not in VERDICTS]
    final = {}
    resolved = Path(args.resolved) if args.resolved else None
    previous = out / "disagreements.csv"
    if resolved is None and previous.exists():
        with open(previous, encoding="utf-8-sig") as f:
            if any(r.get("final", "").strip() for r in csv.DictReader(f)):
                raise ValueError(f"{previous} contains final decisions; supply it with --resolved")
    if resolved:
        with open(resolved, encoding="utf-8-sig") as f:
            seen = set()
            for r in csv.DictReader(f):
                value = normalize_verdict(r.get("final"))
                if not value:
                    continue
                image_id = r["image_id"]
                if image_id in seen:
                    raise ValueError(f"duplicate final decision for {image_id}")
                seen.add(image_id)
                a, b = (got.get((image_id, tab), ("", ""))[0] for tab in TAB_ALIASES)
                if not a or not b or a == b:
                    raise ValueError(f"{image_id}: final decision requires a completed disagreement")
                if (normalize_verdict(r.get("primary")) != a
                        or normalize_verdict(r.get("recheck")) != b):
                    raise ValueError(f"{image_id}: final decision refers to outdated judgments")
                final[image_id] = value
    text = {p["prompt_id"]: p["prompt"] for p in prompts}
    auto = per_challenge_metrics(args.images, key) if args.images else {}

    # Re-check: agreement before the discussion, and the list to discuss.
    disagree, pairs = [], []
    for r in key:
        if not r["recheck_judge"]:
            continue
        (a, note_a), (b, note_b) = (got.get((r["image_id"], tab), ("", "")) for tab in ("Primary", "Recheck"))
        if a in VERDICTS and b in VERDICTS:
            pairs.append((a, b))
        if a in VERDICTS and b in VERDICTS and a != b:
            disagree.append({**r, "prompt": text[r["prompt_id"]], "primary": a, "primary_note": note_a,
                             "recheck": b, "recheck_note": note_b, "final": final.get(r["image_id"], "")})
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "disagreements.csv", "w", newline="", encoding="utf-8-sig") as f:  # BOM: opens in Excel
        fields = ["image_id", "prompt_id", "challenge", "prompt", "judge", "primary", "primary_note",
                  "recheck_judge", "recheck", "recheck_note", "final"]
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(disagree)
    unresolved = [d["image_id"] for d in disagree if not d["final"]]
    agree = sum(a == b for a, b in pairs)
    kappa = cohen_kappa(pairs)

    # Final verdict: the discussed one where there was a disagreement, else the primary one.
    verdict = {r["image_id"]: final.get(r["image_id"]) or got.get((r["image_id"], "Primary"), ("", ""))[0]
               for r in key}
    counts = defaultdict(Counter)
    for r in key:
        for challenge in (r["challenge"], "all"):
            counts[(r["method"], challenge)][verdict[r["image_id"]]] += 1
    status = "final" if not missing and not unresolved else "PROVISIONAL"
    configs = {r["method"]: r["config"] for r in key}
    rows = []
    for method in METHODS:
        for challenge in list(CHALLENGES) + ["all"]:
            c = counts[(method, challenge)]
            judged = c["fail"] + c["pass"]
            lo, hi = wilson(c["fail"], judged)
            a = auto.get((method, challenge), {})
            rows.append({"status": status, "method": METHODS[method], "config": configs[method], "challenge": challenge,
                         "n": sum(c.values()), "fail": c["fail"], "pass": c["pass"], "unsure": c["unsure"],
                         "missing": c[""],
                         "fail_rate": f"{c['fail'] / judged:.3f}" if judged else "",
                         "ci95_low": f"{lo:.3f}" if judged else "", "ci95_high": f"{hi:.3f}" if judged else "",
                         "clip_score": f"{a['clip_score']:.2f}" if a else "",
                         "image_reward": f"{a['image_reward']:.3f}" if a else ""})
    with open(out / "e5_failure.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

    kappa_text = f"{kappa:.2f}" if kappa is not None else "undefined"
    summary = [
        f"status: {status}",
        f"judgments missing or invalid: {len(missing)}" + (f" (e.g. {missing[:5]})" if missing else ""),
        f"re-check pairs: {len(pairs)}; agreement {agree}/{len(pairs)}"
        + (f" ({agree / len(pairs):.1%}), Cohen's kappa {kappa_text}" if pairs else ""),
        f"disagreements: {len(disagree)}, unresolved: {len(unresolved)}",
        "fail_rate = fail / (fail + pass); unsure images are excluded and counted separately;"
        " ci95 is the Wilson score interval",
    ]
    (out / "e5_summary.txt").write_text("\n".join(summary) + "\n", encoding="utf-8")
    print("\n".join(summary))
    plot_failure(rows, out, status)
    if args.packages:
        plot_cases(key, verdict, text, Path(args.packages), out, random.Random(args.seed), status=status)
    else:
        for ext in ("png", "pdf"):
            (out / f"e5_failure_cases.{ext}").unlink(missing_ok=True)
    print(f"written to {out}")


def plot_failure(rows, out, status="PROVISIONAL"):
    """Failure rate with its 95% interval, one panel per challenge."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    colors = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100"]
    markers = ["o", "s", "^", "D"]
    names = list(METHODS.values())
    panels = list(CHALLENGES) + ["all"]
    titles = {"quantity": "Quantity", "writing": "Writing", "position": "Positioning", "all": "All three"}
    fig, axes = plt.subplots(1, 4, figsize=(12.5, 3.3), sharey=True, constrained_layout=True)
    for ax, ch in zip(axes, panels):
        for i, name in enumerate(names):
            r = next(r for r in rows if r["method"] == name and r["challenge"] == ch)
            if not r["fail_rate"]:
                continue
            rate, lo, hi = (float(r[k]) for k in ("fail_rate", "ci95_low", "ci95_high"))
            ax.errorbar(i, rate, yerr=[[rate - lo], [hi - rate]], fmt=markers[i], color=colors[i],
                        markersize=8, markeredgecolor="white", capsize=3, linewidth=1.5)
            ax.annotate(f"{rate:.0%}", (i, rate), xytext=(9, 0), textcoords="offset points",
                        va="center", fontsize=8, color="#52514e")
        ax.set_xticks(range(len(names)), [n.replace("SDXL + DPM-Solver++", "SDXL base") for n in names],
                      rotation=30, ha="right", fontsize=8.5)
        ax.set_xlim(-0.6, len(names) - 0.2)
        ax.set_ylim(0, 1)
        ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
        ax.set_title(titles[ch], loc="left", fontsize=10)
        ax.grid(axis="y", color="#e4e3df")
        ax.set_axisbelow(True)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
    axes[0].set_ylabel("Human-judged failure rate")
    fig.suptitle(f"{status} — E5: failure rate on the selected PartiPrompts subset (bars: 95% Wilson interval)",
                 x=0.01, ha="left", fontsize=11)
    for ext in ("png", "pdf"):
        fig.savefig(out / f"e5_failure_rate.{ext}", dpi=200, facecolor="white")
    plt.close(fig)


def plot_cases(key, verdict, text, packages, out, rng, per_challenge=2, status="PROVISIONAL"):
    """Selected examples: prompts where the methods disagree, all four images side by side.
    The prompts are drawn at random (fixed seed) from those with mixed outcomes."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    # Package folders use English names; the first round's packages carry Chinese names.
    paths = {p.stem: p for pattern in ("E5_judge_*/images/*.png", "E5判定_*/images/*.png")
             for p in packages.glob(pattern)}
    by_prompt = defaultdict(dict)
    for r in key:
        by_prompt[r["prompt_id"]][r["method"]] = r
    chosen = []
    for ch in CHALLENGES:
        mixed = sorted(pid for pid, ms in by_prompt.items()
                       if ms[next(iter(ms))]["challenge"] == ch
                       and {verdict[m["image_id"]] for m in ms.values()} >= {"pass", "fail"})
        chosen += rng.sample(mixed, min(per_challenge, len(mixed)))
    if not chosen:
        for ext in ("png", "pdf"):
            (out / f"e5_failure_cases.{ext}").unlink(missing_ok=True)
        return
    fig = plt.figure(figsize=(10, 3.1 * len(chosen)), layout="constrained")
    label = {"pass": "pass", "fail": "FAIL", "unsure": "unsure", "": "not judged"}
    for sub, pid in zip(fig.subfigures(len(chosen), 1, squeeze=False).flat, chosen):
        rows = by_prompt[pid]
        prompt = text[pid] if len(text[pid]) < 110 else text[pid][:107] + "..."
        sub.suptitle(f"[{rows[next(iter(rows))]['challenge'].capitalize()}] {prompt}",
                     x=0.01, ha="left", fontsize=9)
        for ax, method in zip(sub.subplots(1, 4), METHODS):
            r = rows[method]
            with Image.open(paths[r["image_id"]]) as im:
                ax.imshow(im)
            v = verdict[r["image_id"]]
            ax.set_title(f"{METHODS[method]}: {label[v]}", fontsize=9,
                         color="#c4372c" if v == "fail" else "#0b0b0b")
            ax.set_xticks([])
            ax.set_yticks([])
    fig.suptitle(f"{status} — E5: selected examples (prompts drawn at random among those where the methods "
                 "disagree)", x=0.01, ha="left", fontsize=11)
    for ext in ("png", "pdf"):
        fig.savefig(out / f"e5_failure_cases.{ext}", dpi=150, facecolor="white")
    plt.close(fig)


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("pack")
    p.add_argument("--images", default="outputs/e5/parti")
    p.add_argument("--prompts", default=str(PROMPTS))
    p.add_argument("--out", default="outputs/e5_judging")
    p.add_argument("--seed", type=int, default=SEED)
    s = sub.add_parser("score")
    s.add_argument("--prompts", default=str(PROMPTS))
    s.add_argument("--key", required=True)
    s.add_argument("--sheets", required=True, help="folder with the returned .xlsx sheets")
    s.add_argument("--resolved", help="disagreements.csv with the final column filled in")
    s.add_argument("--out", default="outputs/e5_results")
    s.add_argument("--images", help="E5 output root with per_image_scores.jsonl (auto metrics by challenge)")
    s.add_argument("--packages", help="the judging packages folder, to draw the failure examples")
    s.add_argument("--seed", type=int, default=SEED)
    args = ap.parse_args()
    pack(args) if args.cmd == "pack" else score(args)


if __name__ == "__main__":
    main()
